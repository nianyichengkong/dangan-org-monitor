"""Excel 导出：将当前筛选后的机构列表写入 .xlsx。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADERS = [
    "序号", "行政区划代码", "机构编号", "机构全称", "机构层级",
    "通讯地址", "联系电话", "更新状态", "更新日期",
]
WIDTHS = [6, 12, 12, 42, 10, 52, 26, 10, 12]

# 变更行浅黄、新增行浅绿，与工作台标签色一致
FILL_UPD = PatternFill("solid", fgColor="FFFBE6")
FILL_NEW = PatternFill("solid", fgColor="F6FFED")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def export(rows: list[dict], path: Path) -> None:
    """导出机构列表到 xlsx。

    Args:
        rows: 筛选后的机构字典列表（database.all_orgs 的结构）。
        path: 目标文件路径。

    Raises:
        OSError: 目标路径不可写（如文件被 Excel 占用）。
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "档案管理服务机构"

    # 表头
    for c, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="F2F2F2")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(c)].width = WIDTHS[c - 1]
    ws.freeze_panes = "A2"

    # 数据行
    st_label = {"new": "新增", "upd": "已变更"}
    for i, r in enumerate(rows, 1):
        # 机构全称前缀省份，与工作台表格显示规则一致
        status = st_label.get(r["status"], "—")
        vals = [i, r["admin_code"], r["org_code"],
                f"{r['province']}·{r['name']}", r["level"],
                r["addr"], r["tel"], status, r["updated_at"] or "—"]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=i + 1, column=c, value=v)
            cell.border = BORDER
            if c in (2, 3):
                cell.number_format = "@"  # 代码列按文本，防止丢前导 0
            if r["status"] == "upd":
                cell.fill = FILL_UPD
            elif r["status"] == "new":
                cell.fill = FILL_NEW

    wb.save(path)
